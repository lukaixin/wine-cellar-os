from __future__ import annotations

import hashlib
import io
import json
import re
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Any

from fastapi import HTTPException, status
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import Bottle, Cellar, ImportBatch, ImportBatchRow, Location, Producer, Vintage, Wine
from app.models.base import utcnow
from app.schemas.common import (
    BottleCreate,
    WineListImportConfirmResponse,
    WineListImportDraftResponse,
    WineListImportDraftRow,
    WineListImportNormalizedRow,
    WineListImportRegistryPreview,
)
from app.services.inventory_service import BottleCreationResult, InventoryService
from app.services.registry_service import RegistryService, normalize_registry_text

SHEET_NAME = "Stock"
EXPECTED_HEADERS = [
    "No",
    "Country",
    "Region",
    "Sub-Region",
    "Producer",
    "Wine",
    "Style",
    "Vintage",
    "Quantity",
    "Cost",
    "Channel",
    "WS Price",
    "Status",
    "Total",
    "CAVE",
]
REQUIRED_COLUMNS = ["Country", "Region", "Sub-Region", "Producer", "Style", "Vintage", "Quantity"]
STYLE_MAP = {
    "red": "red",
    "rouge": "red",
    "white": "white",
    "blanc": "white",
    "rose": "rose",
    "rosé": "rose",
    "sweet": "sweet",
    "fortified": "fortified",
    "sparkling": "sparkling",
    "other": "other",
}

LOCATION_CONFIG = {
    "FR_FRIEND_HOME": {
        "cellar_name": "法国朋友家",
        "cellar_type": "other",
        "location_code": "FR",
        "location_type": "freeform",
        "location_label": "法国朋友家 / FR",
    },
    "CAVE_C": {
        "cellar_name": "家中酒柜",
        "cellar_type": "cabinet",
        "location_code": "C",
        "location_type": "bulk",
        "location_label": "家中酒柜 / C",
    },
    "HOME_OUTSIDE_CABINET": {
        "cellar_name": "家中散放",
        "cellar_type": "other",
        "location_code": "HOME-OUTSIDE",
        "location_type": "freeform",
        "location_label": "家中散放 / HOME-OUTSIDE",
    },
}


def _clean(value: object) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value)).replace("\u3000", " ").strip()
    return text or None


def _json_ready(value: object) -> object:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Decimal):
        return str(value)
    return str(value)


def _lookup_key(value: str) -> str:
    text = unicodedata.normalize("NFKC", value).strip().casefold()
    text = text.replace("’", "'").replace("‘", "'").replace("`", "'")
    text = text.replace("‐", "-").replace("‑", "-").replace("–", "-").replace("—", "-")
    text = "".join(ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch))
    text = re.sub(r"[\s\-_./]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


@dataclass(frozen=True)
class ParsedWineListRow:
    row_number: int
    raw: dict[str, Any]
    normalized: WineListImportNormalizedRow | None
    errors: list[str]
    warnings: list[str]
    row_hash: str

    @property
    def status(self) -> str:
        return "invalid" if self.errors else "valid"


@dataclass(frozen=True)
class WineListImportService:
    session: Session

    def create_draft(self, filename: str, content: bytes) -> WineListImportDraftResponse:
        if not content:
            raise HTTPException(status_code=422, detail="WineList file is empty")

        file_sha256 = hashlib.sha256(content).hexdigest()
        workbook = self._load_workbook(content)
        if SHEET_NAME not in workbook.sheetnames:
            raise HTTPException(status_code=422, detail=f"Workbook must include {SHEET_NAME} sheet")

        sheet = workbook[SHEET_NAME]
        headers = [
            _clean(cell.value)
            for cell in next(sheet.iter_rows(min_row=1, max_row=1, max_col=len(EXPECTED_HEADERS)))
        ]
        self._validate_headers(headers)

        parsed_rows = self._parse_rows(sheet, file_sha256)
        if not parsed_rows:
            raise HTTPException(status_code=422, detail="Stock sheet has no inventory rows")

        self._mark_previously_confirmed_rows(parsed_rows)
        registry_preview = self._registry_preview(parsed_rows)
        valid_rows = [row for row in parsed_rows if row.status == "valid" and row.normalized is not None]
        location_counts = self._location_counts(valid_rows)
        warning_summary = sorted(
            {warning for row in parsed_rows for warning in row.warnings if not warning.startswith("Cost is blank")}
        )

        batch = ImportBatch(
            source_filename=filename,
            sheet_name=SHEET_NAME,
            file_sha256=file_sha256,
            status="draft",
            row_count=len(parsed_rows),
            valid_rows=len(valid_rows),
            invalid_rows=len(parsed_rows) - len(valid_rows),
            bottle_count=sum(row.normalized.quantity for row in valid_rows),
            warnings=warning_summary,
        )
        self.session.add(batch)
        self.session.flush()
        db_rows: list[ImportBatchRow] = []
        for row in parsed_rows:
            db_row = ImportBatchRow(
                batch_id=batch.id,
                source_row_number=row.row_number,
                raw_row=row.raw,
                normalized_row=row.normalized.model_dump(mode="json") if row.normalized is not None else None,
                row_hash=row.row_hash,
                status=row.status,
                errors=row.errors,
                warnings=row.warnings,
                created_bottle_ids=[],
                created_transaction_ids=[],
            )
            self.session.add(db_row)
            db_rows.append(db_row)
        self.session.flush()
        self._record_registry_mappings(batch, db_rows)
        self.session.commit()
        self.session.refresh(batch)
        for db_row in db_rows:
            self.session.refresh(db_row)

        return self._draft_response(batch, db_rows, registry_preview, location_counts)

    def confirm(self, batch_id: str) -> WineListImportConfirmResponse:
        try:
            batch = self.session.get(ImportBatch, batch_id)
            if batch is None:
                raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="WineList import batch not found")
            if batch.status == "confirmed":
                raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="WineList import batch already confirmed")
            if batch.status != "draft":
                raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=f"Batch status is {batch.status}")

            rows = self.session.scalars(
                select(ImportBatchRow)
                .where(ImportBatchRow.batch_id == batch.id)
                .order_by(ImportBatchRow.source_row_number)
            ).all()
            invalid_rows = [row for row in rows if row.status != "valid" or not row.normalized_row]
            if invalid_rows:
                raise HTTPException(status_code=422, detail="Cannot confirm WineList import with invalid rows")

            duplicate_hashes = self._confirmed_row_hashes([row.row_hash for row in rows], exclude_batch_id=batch.id)
            if duplicate_hashes:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail="WineList rows were already confirmed in another batch",
                )

            normalized_rows = [(row, WineListImportNormalizedRow.model_validate(row.normalized_row)) for row in rows]
            location_counts = self._location_counts(
                [ParsedWineListRow(row.source_row_number, row.raw_row, normalized, [], [], row.row_hash) for row, normalized in normalized_rows]
            )
            locations = self._resolve_or_create_locations(location_counts)

            created_bottle_ids: list[str] = []
            created_transaction_ids: list[str] = []
            warnings = list(batch.warnings or [])
            for db_row, normalized in normalized_rows:
                row_bottle_ids: list[str] = []
                row_transaction_ids: list[str] = []
                location = locations[normalized.location_key]
                for bottle_index in range(1, normalized.quantity + 1):
                    result = self._create_imported_bottle(normalized, location, batch, db_row, bottle_index)
                    row_bottle_ids.append(result.bottle.id)
                    row_transaction_ids.append(result.transaction.id)
                    created_bottle_ids.append(result.bottle.id)
                    created_transaction_ids.append(result.transaction.id)
                db_row.status = "imported"
                db_row.created_bottle_ids = row_bottle_ids
                db_row.created_transaction_ids = row_transaction_ids
                self.session.add(db_row)

            batch.status = "confirmed"
            batch.confirmed_at = utcnow()
            batch.bottle_count = len(created_bottle_ids)
            batch.warnings = sorted(set(warnings))
            self.session.add(batch)
            self.session.commit()

            return WineListImportConfirmResponse(
                batch_id=batch.id,
                created_bottles=len(created_bottle_ids),
                created_bottle_ids=created_bottle_ids,
                created_transaction_ids=created_transaction_ids,
                location_counts=location_counts,
                warnings=batch.warnings or [],
            )
        except HTTPException:
            self.session.rollback()
            raise
        except Exception as exc:
            self.session.rollback()
            raise HTTPException(status_code=500, detail=f"WineList import confirm rolled back: {exc}") from exc

    def _load_workbook(self, content: bytes):
        try:
            return load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            try:
                return load_workbook(self._workbook_without_styles(content), read_only=True, data_only=True)
            except Exception:
                raise HTTPException(status_code=422, detail=f"Cannot read WineList workbook: {exc}") from exc

    def _workbook_without_styles(self, content: bytes) -> io.BytesIO:
        sanitized = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(
            sanitized,
            mode="w",
            compression=zipfile.ZIP_DEFLATED,
        ) as target:
            for item in source.infolist():
                if item.filename == "xl/styles.xml":
                    continue
                target.writestr(item, source.read(item.filename))
        sanitized.seek(0)
        return sanitized

    def _validate_headers(self, headers: list[str | None]) -> None:
        actual = [header for header in headers[: len(EXPECTED_HEADERS)]]
        missing = [header for header in EXPECTED_HEADERS if header not in actual]
        if missing:
            raise HTTPException(status_code=422, detail=f"Stock sheet missing headers: {', '.join(missing)}")

    def _parse_rows(self, sheet, file_sha256: str) -> list[ParsedWineListRow]:
        rows: list[ParsedWineListRow] = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, max_col=len(EXPECTED_HEADERS), values_only=True),
            start=2,
        ):
            raw = {header: _json_ready(value) for header, value in zip(EXPECTED_HEADERS, values, strict=False)}
            if not any(_clean(raw.get(field)) for field in ["Producer", "Wine", "Vintage", "Quantity"]):
                continue
            errors: list[str] = []
            warnings: list[str] = []
            normalized = self._normalize_row(raw, errors, warnings)
            row_hash = self._row_hash(file_sha256, row_number, raw)
            rows.append(
                ParsedWineListRow(
                    row_number=row_number,
                    raw=raw,
                    normalized=normalized,
                    errors=errors,
                    warnings=warnings,
                    row_hash=row_hash,
                )
            )
        return rows

    def _normalize_row(
        self,
        raw: dict[str, Any],
        errors: list[str],
        warnings: list[str],
    ) -> WineListImportNormalizedRow | None:
        missing_required = [field for field in REQUIRED_COLUMNS if not _clean(raw.get(field))]
        errors.extend(f"{field} is required" for field in missing_required)

        country = _clean(raw.get("Country"))
        region = self._region(_clean(raw.get("Region")), warnings)
        appellation = _clean(raw.get("Sub-Region"))
        producer_name = _clean(raw.get("Producer"))
        wine_name = _clean(raw.get("Wine"))
        if wine_name is None and producer_name is not None:
            wine_name = producer_name
            warnings.append("Wine is blank; using Producer as wine name")
        wine_type = self._style(_clean(raw.get("Style")), errors)
        vintage_year = self._vintage_year(_clean(raw.get("Vintage")), errors)
        quantity = self._quantity(_clean(raw.get("Quantity")), errors)
        purchase_price = self._cost(_clean(raw.get("Cost")), errors, warnings)
        ws_price = self._ws_price(_clean(raw.get("WS Price")), errors)
        purchase_channel = _clean(raw.get("Channel"))
        raw_status = _clean(raw.get("Status"))
        raw_cave = _clean(raw.get("CAVE"))
        location_key = self._location_key(raw_status, raw_cave, errors)

        if errors:
            return None
        assert country is not None
        assert region is not None
        assert appellation is not None
        assert producer_name is not None
        assert wine_name is not None
        assert wine_type is not None
        assert vintage_year is not None
        assert quantity is not None
        assert location_key is not None
        location = LOCATION_CONFIG[location_key]
        registry_decision = RegistryService(self.session).evaluate_winelist_row(
            producer_name=producer_name,
            wine_name=wine_name,
            vintage_year=vintage_year,
            country=country,
            region=region,
            appellation=appellation,
            wine_type=wine_type,
        )
        return WineListImportNormalizedRow(
            country=country,
            region=region,
            appellation=appellation,
            producer_name=producer_name,
            wine_name=wine_name,
            type=wine_type,
            vintage_year=vintage_year,
            quantity=quantity,
            bottle_size_ml=750,
            purchase_price=purchase_price,
            purchase_currency="CNY",
            purchase_channel=purchase_channel,
            purchase_source="WineList.xlsx Stock",
            ws_price=ws_price,
            cellar_name=location["cellar_name"],
            location_code=location["location_code"],
            location_label=location["location_label"],
            location_key=location_key,
            raw_status=raw_status,
            raw_cave=raw_cave,
            producer_match_score=registry_decision.producer_match.match_score,
            wine_match_score=registry_decision.wine_match.match_score,
            vintage_match_score=registry_decision.vintage_match.match_score,
            overall_match_score=registry_decision.overall_match_score,
            decision=registry_decision.decision,
        )

    def _region(self, value: str | None, warnings: list[str]) -> str | None:
        if value is None:
            return None
        if _lookup_key(value) == "sorthern rhone":
            warnings.append("Region normalized from Sorthern Rhone to Southern Rhone")
            return "Southern Rhone"
        return value

    def _style(self, value: str | None, errors: list[str]) -> str | None:
        if value is None:
            return None
        mapped = STYLE_MAP.get(_lookup_key(value))
        if mapped is None:
            errors.append(f"Style cannot be mapped: {value}")
        return mapped

    def _vintage_year(self, value: str | None, errors: list[str]) -> int | None:
        if value is None:
            return None
        try:
            year = int(Decimal(value))
        except (InvalidOperation, ValueError):
            errors.append(f"Vintage must be a year: {value}")
            return None
        if year < 1800 or year > 2100:
            errors.append(f"Vintage year is out of range: {year}")
            return None
        return year

    def _quantity(self, value: str | None, errors: list[str]) -> int | None:
        if value is None:
            return None
        try:
            parsed = Decimal(value)
        except InvalidOperation:
            errors.append(f"Quantity must be a positive integer: {value}")
            return None
        if parsed != parsed.to_integral_value() or parsed <= 0:
            errors.append(f"Quantity must be a positive integer: {value}")
            return None
        return int(parsed)

    def _cost(self, value: str | None, errors: list[str], warnings: list[str]) -> Decimal | None:
        if value is None:
            warnings.append("Cost is blank; purchase_price will be null")
            return None
        try:
            cost = Decimal(value).quantize(Decimal("0.01"))
        except InvalidOperation:
            errors.append(f"Cost must be a non-negative decimal: {value}")
            return None
        if cost < 0:
            errors.append(f"Cost must be a non-negative decimal: {value}")
            return None
        return cost

    def _ws_price(self, value: str | None, errors: list[str]) -> Decimal | None:
        if value is None:
            return None
        try:
            price = Decimal(value).quantize(Decimal("0.01"))
        except InvalidOperation:
            errors.append(f"WS Price must be a non-negative decimal: {value}")
            return None
        if price < 0:
            errors.append(f"WS Price must be a non-negative decimal: {value}")
            return None
        return price

    def _location_key(self, raw_status: str | None, raw_cave: str | None, errors: list[str]) -> str | None:
        status_key = _lookup_key(raw_status) if raw_status else ""
        cave_key = _lookup_key(raw_cave) if raw_cave else ""
        if status_key == "fr":
            return "FR_FRIEND_HOME"
        if cave_key == "c":
            return "CAVE_C"
        if not status_key and not cave_key:
            return "HOME_OUTSIDE_CABINET"
        errors.append(f"Cannot map location from Status={raw_status or ''}, CAVE={raw_cave or ''}")
        return None

    def _row_hash(self, file_sha256: str, row_number: int, raw: dict[str, Any]) -> str:
        payload = {
            "file_sha256": file_sha256,
            "sheet": SHEET_NAME,
            "source_row_number": row_number,
            "raw": raw,
        }
        encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")
        return hashlib.sha256(encoded).hexdigest()

    def _mark_previously_confirmed_rows(self, rows: list[ParsedWineListRow]) -> None:
        confirmed = self._confirmed_row_hashes([row.row_hash for row in rows])
        if not confirmed:
            return
        for row in rows:
            if row.row_hash in confirmed:
                row.errors.append("Row was already confirmed in a previous WineList import batch")

    def _confirmed_row_hashes(self, row_hashes: list[str], exclude_batch_id: str | None = None) -> set[str]:
        if not row_hashes:
            return set()
        query = (
            select(ImportBatchRow.row_hash)
            .join(ImportBatch, ImportBatch.id == ImportBatchRow.batch_id)
            .where(ImportBatch.status == "confirmed", ImportBatchRow.row_hash.in_(row_hashes))
        )
        if exclude_batch_id is not None:
            query = query.where(ImportBatch.id != exclude_batch_id)
        return set(self.session.scalars(query).all())

    def _record_registry_mappings(self, batch: ImportBatch, rows: list[ImportBatchRow]) -> None:
        registry = RegistryService(self.session)
        for row in rows:
            if row.normalized_row is None:
                if any("is required" in error for error in row.errors or []):
                    registry.enqueue_review(
                        review_type="missing_field",
                        title=f"WineList row {row.source_row_number} has missing fields",
                        detail="Draft row is invalid and requires manual source cleanup before confirm.",
                        source="winelist:Stock",
                        raw_value=row.raw_row,
                        evidence={"errors": row.errors or []},
                    )
                continue

            normalized = WineListImportNormalizedRow.model_validate(row.normalized_row)
            registry_decision = registry.evaluate_winelist_row(
                producer_name=normalized.producer_name,
                wine_name=normalized.wine_name,
                vintage_year=normalized.vintage_year,
                country=normalized.country,
                region=normalized.region,
                appellation=normalized.appellation,
                wine_type=normalized.type,
            )
            matches = [
                ("Producer", normalized.producer_name, registry_decision.producer_match),
                ("Wine", normalized.wine_name, registry_decision.wine_match),
                ("Vintage", str(normalized.vintage_year), registry_decision.vintage_match),
            ]
            for field_name, raw_value, match in matches:
                mapping = registry.record_import_mapping(
                    source="winelist:Stock",
                    import_batch_id=batch.id,
                    import_batch_row_id=row.id,
                    source_row_number=row.source_row_number,
                    raw_value={"field": field_name, "value": raw_value},
                    match=match,
                )
                if match.matched_by == "alias_conflict":
                    registry.enqueue_review(
                        review_type="alias_conflict",
                        title=f"WineList row {row.source_row_number} alias conflict",
                        detail=f"{field_name} alias points to more than one registry entity.",
                        source="winelist:Stock",
                        entity_type=match.entity_type,
                        entity_id=match.target_id,
                        import_mapping_id=mapping.id,
                        raw_value={"field": field_name, "value": raw_value},
                        match_score=match.match_score,
                        evidence=match.evidence,
                    )
                elif match.decision == "needs_review":
                    registry.enqueue_review(
                        review_type="match_range",
                        title=f"WineList row {row.source_row_number} needs registry review",
                        detail=f"{field_name} matched at {match.match_score}; review before treating it as stable.",
                        source="winelist:Stock",
                        entity_type=match.entity_type,
                        entity_id=match.target_id,
                        import_mapping_id=mapping.id,
                        raw_value={"field": field_name, "value": raw_value},
                        match_score=match.match_score,
                        evidence=match.evidence,
                    )

    def _registry_preview(self, rows: list[ParsedWineListRow]) -> WineListImportRegistryPreview:
        valid = [row.normalized for row in rows if row.status == "valid" and row.normalized is not None]
        registry = RegistryService(self.session)

        seen_producers: set[str] = set()
        seen_wines: set[tuple[str, str]] = set()
        seen_vintages: set[tuple[str, str, int]] = set()
        producer_existing = 0
        producer_new = 0
        wine_existing = 0
        wine_new = 0
        vintage_existing = 0
        vintage_new = 0

        for row in valid:
            producer_key = normalize_registry_text(row.producer_name)
            producer_match = registry.match_producer(row.producer_name)
            if producer_key not in seen_producers:
                seen_producers.add(producer_key)
                if producer_match.target_id is not None and producer_match.decision == "auto_accepted":
                    producer_existing += 1
                else:
                    producer_new += 1

            producer = self.session.get(Producer, producer_match.target_id) if producer_match.target_id else None
            wine_match = registry.match_wine(
                producer,
                row.wine_name,
                country=row.country,
                region=row.region,
                appellation=row.appellation,
                wine_type=row.type,
            )
            wine_key = (producer_key, normalize_registry_text(row.wine_name))
            if wine_key not in seen_wines:
                seen_wines.add(wine_key)
                if wine_match.target_id is not None and wine_match.decision == "auto_accepted":
                    wine_existing += 1
                else:
                    wine_new += 1

            wine = self.session.get(Wine, wine_match.target_id) if wine_match.target_id else None
            vintage_match = registry.match_vintage(wine, row.vintage_year)
            vintage_key = (producer_key, normalize_registry_text(row.wine_name), row.vintage_year)
            if vintage_key not in seen_vintages:
                seen_vintages.add(vintage_key)
                if vintage_match.target_id is not None and vintage_match.decision == "auto_accepted":
                    vintage_existing += 1
                else:
                    vintage_new += 1

        return WineListImportRegistryPreview(
            producers_new=producer_new,
            producers_existing=producer_existing,
            wines_new=wine_new,
            wines_existing=wine_existing,
            vintages_new=vintage_new,
            vintages_existing=vintage_existing,
        )

    def _location_counts(self, rows: list[ParsedWineListRow]) -> dict[str, int]:
        counts: Counter[str] = Counter()
        for row in rows:
            if row.normalized is None:
                continue
            counts[row.normalized.location_label] += row.normalized.quantity
        return dict(counts)

    def _resolve_or_create_locations(self, location_counts: dict[str, int]) -> dict[str, Location]:
        locations: dict[str, Location] = {}
        for location_key, config in LOCATION_CONFIG.items():
            count = location_counts.get(config["location_label"], 0)
            if count <= 0:
                continue
            cellar = self._get_or_create_cellar(config["cellar_name"], config["cellar_type"])
            location = self._get_or_create_location(cellar, config, count)
            occupied = self._occupied_count(location.id)
            if occupied + count > location.capacity:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail=(
                        f"Location {cellar.name}/{location.code} capacity is too small "
                        f"({occupied}+{count}/{location.capacity}); expand it before confirming"
                    ),
                )
            locations[location_key] = location
        return locations

    def _get_or_create_cellar(self, name: str, cellar_type: str) -> Cellar:
        cellar = self.session.scalar(select(Cellar).where(func.lower(Cellar.name) == name.lower()))
        if cellar is not None:
            return cellar
        cellar = Cellar(name=name, cellar_type=cellar_type, notes="Created by WineList import")
        self.session.add(cellar)
        self.session.flush()
        return cellar

    def _get_or_create_location(self, cellar: Cellar, config: dict[str, str], required_capacity: int) -> Location:
        location = self.session.scalar(
            select(Location).where(
                Location.cellar_id == cellar.id,
                func.lower(Location.code) == config["location_code"].lower(),
            )
        )
        if location is not None:
            return location
        location = Location(
            cellar_id=cellar.id,
            code=config["location_code"],
            capacity=required_capacity,
            location_type=config["location_type"],
            position_note="WineList import bulk location",
        )
        self.session.add(location)
        self.session.flush()
        return location

    def _occupied_count(self, location_id: str) -> int:
        return self.session.scalar(
            select(func.count(Bottle.id)).where(
                Bottle.current_location_id == location_id,
                Bottle.current_status == "in_stock",
            )
        ) or 0

    def _create_imported_bottle(
        self,
        row: WineListImportNormalizedRow,
        location: Location,
        batch: ImportBatch,
        db_row: ImportBatchRow,
        bottle_index: int,
    ) -> BottleCreationResult:
        producer = self._get_or_create_producer(row)
        wine = self._get_or_create_wine(row, producer)
        vintage = self._get_or_create_vintage(row, wine)
        self.session.flush()
        return InventoryService(self.session).create_bottle_with_in_transaction_result(
            BottleCreate(
                vintage_id=vintage.id,
                current_location_id=location.id,
                bottle_size_ml=row.bottle_size_ml,
                purchase_price=row.purchase_price,
                purchase_currency=row.purchase_currency,
                purchase_source=f"{batch.source_filename}:{SHEET_NAME}:row {db_row.source_row_number}",
                purchase_channel=row.purchase_channel,
                ws_price=row.ws_price,
                personal_notes=f"WineList import batch {batch.id} row {db_row.source_row_number}",
                actor="winelist_import",
                note=(
                    f"WineList import batch {batch.id} row {db_row.source_row_number} "
                    f"bottle {bottle_index}/{row.quantity}"
                ),
            ),
            commit=False,
        )

    def _get_or_create_producer(self, row: WineListImportNormalizedRow) -> Producer:
        registry = RegistryService(self.session)
        producer_match = registry.match_producer(row.producer_name)
        producer = (
            self.session.get(Producer, producer_match.target_id)
            if producer_match.target_id is not None and producer_match.decision == "auto_accepted"
            else None
        )
        if producer is not None:
            producer.country = producer.country or row.country
            producer.region = producer.region or row.region
            registry.apply_producer_fields(producer)
            self.session.add(producer)
            return producer
        producer = Producer(
            name_original=row.producer_name,
            country=row.country,
            region=row.region,
            aliases=None,
        )
        registry.apply_producer_fields(producer)
        self.session.add(producer)
        self.session.flush()
        return producer

    def _get_or_create_wine(self, row: WineListImportNormalizedRow, producer: Producer) -> Wine:
        registry = RegistryService(self.session)
        wine_match = registry.match_wine(
            producer,
            row.wine_name,
            country=row.country,
            region=row.region,
            appellation=row.appellation,
            wine_type=row.type,
        )
        wine = (
            self.session.get(Wine, wine_match.target_id)
            if wine_match.target_id is not None and wine_match.decision == "auto_accepted"
            else None
        )
        if wine is not None:
            wine.country = wine.country or row.country
            wine.region = wine.region or row.region
            wine.appellation = wine.appellation or row.appellation
            wine.type = wine.type or row.type
            registry.apply_wine_fields(wine, producer)
            self.session.add(wine)
            return wine
        wine = Wine(
            producer_id=producer.id,
            name_original=row.wine_name,
            type=row.type,
            country=row.country,
            region=row.region,
            appellation=row.appellation,
            aliases=None,
        )
        registry.apply_wine_fields(wine, producer)
        self.session.add(wine)
        self.session.flush()
        return wine

    def _get_or_create_vintage(self, row: WineListImportNormalizedRow, wine: Wine) -> Vintage:
        vintage = self.session.scalar(
            select(Vintage).where(
                Vintage.wine_id == wine.id,
                Vintage.vintage_year == row.vintage_year,
                Vintage.is_non_vintage.is_(False),
            )
        )
        if vintage is not None:
            RegistryService(self.session).apply_vintage_fields(vintage)
            return vintage
        vintage = Vintage(wine_id=wine.id, vintage_year=row.vintage_year, is_non_vintage=False)
        RegistryService(self.session).apply_vintage_fields(vintage)
        self.session.add(vintage)
        self.session.flush()
        return vintage

    def _draft_response(
        self,
        batch: ImportBatch,
        rows: list[ImportBatchRow],
        registry_preview: WineListImportRegistryPreview,
        location_counts: dict[str, int],
    ) -> WineListImportDraftResponse:
        return WineListImportDraftResponse(
            batch_id=batch.id,
            filename=batch.source_filename,
            sheet_name=batch.sheet_name,
            file_sha256=batch.file_sha256,
            total_rows=batch.row_count,
            valid_rows=batch.valid_rows,
            invalid_rows=batch.invalid_rows,
            bottle_count=batch.bottle_count,
            location_counts=location_counts,
            registry_preview=registry_preview,
            warnings=batch.warnings or [],
            rows=[
                WineListImportDraftRow(
                    id=row.id,
                    row_number=row.source_row_number,
                    status=row.status if row.status in {"valid", "invalid"} else "valid",
                    raw=row.raw_row,
                    normalized=(
                        WineListImportNormalizedRow.model_validate(row.normalized_row)
                        if row.normalized_row is not None
                        else None
                    ),
                    errors=row.errors or [],
                    warnings=row.warnings or [],
                    bottle_count=(
                        WineListImportNormalizedRow.model_validate(row.normalized_row).quantity
                        if row.normalized_row is not None and row.status == "valid"
                        else 0
                    ),
                )
                for row in rows
            ],
        )
